"""Запускать в папке kaspi-hub:  python3 patch_pricelist.py
Переделывает «Прайс-лист» (показы, CTR, CR, реклама, ДРР) и добавляет статистику рекламы в «Маркетинг». Копии: *.bak3"""
import re, ast, shutil, sys

SERVER_PRICELIST = r'''def pricelist(p):
    a, b = rng(p)
    g, S = getset(), load_stock()
    C = {r[0]: r for r in q("select sku,cost,delivery,mkt,comm from costs")}
    z = lambda: {"sold": 0, "revenue": 0, "ret": 0, "spend": 0, "shows": 0, "clicks": 0, "aord": 0}
    agg = {x["sku"]: z() for x in S}
    for name, code, qty, total, ret in sold(a, b):
        x = match(code, name, S)
        if x:
            r = agg[x["sku"]]
            r["sold"] += qty
            r["revenue"] += total
            r["ret"] += ret
    day = lambda t: datetime.fromtimestamp(t / 1000, TZ).strftime("%Y-%m-%d")
    tot = z()
    for nm, sk, sp, sh, cl, od in q("select name,sku,spend,shows,clicks,orders from ads where day between ? and ?", (day(a), day(b))):
        v = {"spend": sp or 0, "shows": sh or 0, "clicks": cl or 0, "aord": od or 0}
        for k in v:
            tot[k] += v[k]
        key = None
        if sk:
            x = match(str(sk), "", S)
            key = x["sku"] if x else None
        if key is None:  # реклама привязывается к товару по SKU или если в названии есть SKU/название товара
            n = (nm or "").lower()
            for x in S:
                if x["sku"].lower() in n or x["name"].lower() in n:
                    key = x["sku"]
                    break
        if key:
            for k in v:
                agg[key][k] += v[k]
    rows = []
    for x in S:
        _, cost, od, om, oc = C.get(x["sku"], (None,) * 5)
        v = agg[x["sku"]]
        cm, dl = (g["comm"] if oc is None else oc), (g["delivery"] if od is None else od)
        rev, n = v["revenue"], v["sold"]
        profit = None if cost is None else rev * (1 - (cm + g["tax"]) / 100) - n * (cost + dl)
        ad = v["spend"]
        net = None if profit is None else profit - ad
        rows.append({"sku": x["sku"], "name": x["name"], "price": x["price"], "units": x["units"], "sold": n, "revenue": rev,
                     "cost": cost, "cost_total": None if cost is None else n * cost, "profit": profit,
                     "margin": profit / rev * 100 if profit is not None and rev else None,
                     "ret": v["ret"], "ret_pct": v["ret"] / (n + v["ret"]) * 100 if n + v["ret"] else None,
                     "shows": v["shows"] or None, "ctr": v["clicks"] / v["shows"] * 100 if v["shows"] else None,
                     "cr": v["aord"] / v["clicks"] * 100 if v["clicks"] and v["aord"] else None,
                     "ads": ad, "drr": ad / rev * 100 if ad and rev else None, "net": net,
                     "net_margin": net / rev * 100 if net is not None and rev else None})
    kn = [r for r in rows if r["profit"] is not None]
    rk, pr = sum(r["revenue"] for r in kn), sum(r["profit"] for r in kn)
    k = {"sold": sum(r["sold"] for r in rows), "revenue": sum(r["revenue"] for r in rows), "cost": sum(r["cost_total"] for r in kn),
         "profit": pr, "margin": pr / rk * 100 if rk else None, "ads": tot["spend"], "net": pr - tot["spend"],
         "net_margin": (pr - tot["spend"]) / rk * 100 if rk else None,
         "ctr": tot["clicks"] / tot["shows"] * 100 if tot["shows"] else None,
         "cr": tot["aord"] / tot["clicks"] * 100 if tot["clicks"] and tot["aord"] else None,
         "no_cost": sum(1 for r in rows if r["cost"] is None and r["sold"])}
    return {"settings": g, "rows": rows, "k": k}


_n = {}


def import_ads(fname, b64):
    """Отчёт из рекламного кабинета (xlsx/csv): колонки ищутся по названию: дата, SKU/товар, показы, клики, заказы, расход."""
    import io, csv
    raw = base64.b64decode(b64)
    try:
        if fname.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            wb.close()
        else:
            txt = raw.decode("utf-8-sig", errors="replace")
            rows = list(csv.reader(io.StringIO(txt), delimiter=";" if txt.count(";") > txt.count(",") else ","))
    except Exception as e:
        raise ValueError(f"Не удалось прочитать файл: {e}")
    syn = {"day": ("дата", "date", "день"), "sku": ("sku", "артикул", "код"), "name": ("товар", "название", "наименование", "кампания", "модель", "product", "name"),
           "shows": ("показ", "просмотр", "impression", "shows"), "clicks": ("клик", "переход", "click"), "orders": ("заказ", "order", "покупк"),
           "spend": ("расход", "затрат", "потрачен", "стоимост", "бюджет", "spend", "cost", "сумма")}
    hdr = None
    for i, r in enumerate(rows[:15]):
        cells = [str(c or "").strip().lower() for c in r]
        m = {k: next((j for j, c in enumerate(cells) if any(c.startswith(s) for s in v)), None) for k, v in syn.items()}
        if m["spend"] is not None or m["shows"] is not None:
            hdr = (i, m)
            break
    if not hdr:
        raise ValueError("Не нашёл колонки: нужны хотя бы «Показы» или «Расход»")
    i0, m = hdr

    def num(x):
        try:
            return float(str(x).replace("\xa0", "").replace(" ", "").replace(",", ".")) if x not in (None, "") else 0.0
        except ValueError:
            return 0.0

    def dt(x):
        if hasattr(x, "strftime"):
            return x.strftime("%Y-%m-%d")
        s = str(x or "").strip()[:10]
        for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(s, f).strftime("%Y-%m-%d")
            except ValueError:
                pass
        return datetime.now(TZ).strftime("%Y-%m-%d")
    cnt = 0
    for r in rows[i0 + 1:]:
        g = lambda k: r[m[k]] if m[k] is not None and m[k] < len(r) else None
        sku, nm = str(g("sku") or "").strip(), str(g("name") or "").strip()
        if not (sku or nm) or str(sku + nm).lower().startswith(("итого", "total")):
            continue
        d = dt(g("day"))
        db.execute("delete from ads where day=? and coalesce(sku,name)=?", (d, sku or nm[:80]))
        db.execute("insert into ads(day,name,spend,sku,shows,clicks,orders) values(?,?,?,?,?,?,?)",
                   (d, (nm or sku)[:80], num(g("spend")), sku or None, int(num(g("shows"))), int(num(g("clicks"))), int(num(g("orders")))))
        cnt += 1
    if not cnt:
        raise ValueError("В файле нет строк с товарами")
    return cnt
'''
POST_NEW = r'''        elif path == "/api/adstat":
            datetime.fromisoformat(d["day"])
            sk = str(d.get("sku") or "").strip()
            nm = str(d.get("name") or sk or "Реклама")[:80]
            db.execute("insert into ads(day,name,spend,sku,shows,clicks,orders) values(?,?,?,?,?,?,?)",
                       (d["day"], nm, float(d.get("spend") or 0), sk or None, int(float(d.get("shows") or 0)), int(float(d.get("clicks") or 0)), int(float(d.get("orders") or 0))))
            _log("Маркетинг", f"Статистика рекламы {nm}, {d['day']}: показы {d.get('shows') or 0}, клики {d.get('clicks') or 0}, расход {d.get('spend') or 0} ₸")
        elif path == "/api/ad/import":
            _n["n"] = import_ads(str(d["name"]), d["data"])
            _log("Маркетинг", f"Загружен отчёт по рекламе «{d['name']}»: строк {_n['n']}")
        elif path == "/api/syncnow":'''

JS_PRICE = r'''let PR={from:ago(29),to:iso(new Date())},PL={k:"revenue",d:-1,pg:0,sz:10,q:""};
async function pricelist(){
  let j=await api("/api/pricelist",PR),d=j.data;
  const sets=[["comm","Комиссия Kaspi, %"],["delivery","Доставка, ₸ за шт."],["tax","Налог, %"]],pc=x=>x==null?"—":x.toFixed(1).replace(".",",")+"%",cap=+(P.drr_cap)||7;
  const pb=(l,n)=>`<button data-pd="${n}" class="${PR.from===ago(n-1)&&PR.to===iso(new Date())?"on":""}">${l}</button>`;
  const cols=[["name","Товар","Название и SKU товара"],["price","Цена","Цена на Kaspi из ACTIVE.xlsx"],["units","Остаток","Остаток из ACTIVE.xlsx"],["sold","Продано, шт","Сколько штук продано за период без возвратов"],["revenue","Выручка","Выручка за период без отмен и возвратов"],["cost_total","Себестоимость","Себестоимость проданных штук. Ниже вводится цена закупки за 1 шт."],["profit","Прибыль","Выручка минус комиссия, налог, доставка и себестоимость. Без рекламы."],["margin","Маржа","Прибыль в процентах от выручки"],["ret","Возвраты","Сколько штук вернули и какая это доля от проданного"],["shows","Показы","Сколько раз товар показался в рекламе Kaspi."],["ctr","CTR","Сколько из ста показов рекламы закончились переходом на карточку. Цель: не меньше 3,5%."],["cr","CR","Сколько из ста переходов на карточку закончились заказом. Цель: не меньше 2,5%."],["ads","Реклама","Сколько тенге потрачено на рекламу этого товара за период."],["drr","ДРР","Доля расходов на рекламу: сколько тенге рекламы пришлось на каждые сто тенге выручки. Чем меньше, тем лучше. Потолок: "+cap+"%."],["net_margin","Чистая маржа","Маржа с учётом рекламы"]];
  $("#view").innerHTML=head("Прайс-лист",j," · цены и остатки из ACTIVE.xlsx, продажи из Kaspi, реклама из раздела «Маркетинг»")+`
  <div class="bar">${pb("1 нед.",7)}${pb("1 мес.",30)}${pb("3 мес.",90)}${pb("1 год",365)}<input type="date" id="f" value="${PR.from}"><input type="date" id="t" value="${PR.to}">
  <input id="s" placeholder="Поиск по названию или SKU" value="${esc(PL.q)}" style="min-width:240px"><button id="rf">Обновить</button><button id="ex">Выгрузить в файл</button></div>
  <div class="grid" id="kp"></div>
  <div class="panel"><h2>Общие ставки <span class="mut" style="font-weight:400">· нужны для расчёта прибыли</span></h2><div class="bar" style="margin:0">${sets.map(([k,l])=>`<label class="mut">${l}<br><input class="in s" data-s="${k}" type="number" step="any" value="${d.settings[k]}"></label>`).join("")}</div></div>
  <div class="panel tw" style="margin-top:12px"><table><thead><tr>${cols.map(([k,l,t],i)=>`<th data-o="${k}" title="${esc(t)}" class="${i?"r":""}" style="cursor:pointer"><span style="border-bottom:1px dotted var(--mute)">${l}</span>${PL.k===k?(PL.d>0?" ▲":" ▼"):""}</th>`).join("")}</tr></thead><tbody id="tb"></tbody></table></div><div class="bar" id="pgb"></div>`;
  const view=()=>{const s=PL.q.toLowerCase(),k=PL.k;
    return d.rows.filter(r=>(r.name+" "+r.sku).toLowerCase().includes(s)).sort((a,b)=>{const x=a[k],y=b[k];if(x==null&&y==null)return 0;if(x==null)return 1;if(y==null)return -1;return (typeof x=="string"?x.localeCompare(y):x-y)*PL.d})};
  const amb=(v,goal)=>v==null?"—":`<span style="color:${v>=goal?"var(--ok)":"var(--amber)"}">${pc(v)}</span>`;
  const draw=()=>{const K=d.k,R=view(),pages=Math.max(1,Math.ceil(R.length/PL.sz));if(PL.pg>=pages)PL.pg=pages-1;
    $("#kp").innerHTML=card("Продано, шт",f(K.sold))+card("Выручка",m(K.revenue),"ok")+card("Себестоимость",m(K.cost))+card("Прибыль",m(K.profit),K.profit<0?"bad":"ok")+card("Маржа",pc(K.margin))+card("Чистая маржа",pc(K.net_margin),K.net<0?"bad":"")+card("CTR",K.ctr==null?"—":`<span style="color:${K.ctr>=3.5?"var(--ok)":"var(--amber)"}">${pc(K.ctr)}</span>`)+card("CR",K.cr==null?"—":`<span style="color:${K.cr>=2.5?"var(--ok)":"var(--amber)"}">${pc(K.cr)}</span>`)+(K.no_cost?card("Без себестоимости",f(K.no_cost),"warn"):"");
    $("#tb").innerHTML=R.slice(PL.pg*PL.sz,(PL.pg+1)*PL.sz).map(r=>`<tr><td>${esc(r.name)}<div class="mut" style="font-size:12px">${esc(r.sku)}</div></td><td class="r">${f(r.price)} ₸</td><td class="r ${r.units<=0?"red":""}">${f(r.units)}</td><td class="r">${f(r.sold)}</td><td class="r">${f(r.revenue)} ₸</td>
    <td class="r">${r.cost_total==null?"—":f(r.cost_total)+" ₸"}<br><input class="in" data-sku="${esc(r.sku)}" data-k="cost" type="number" step="any" value="${r.cost??""}" placeholder="за шт."></td>
    <td class="r ${r.profit==null?"mut":r.profit<0?"red":"grn"}">${r.profit==null?"—":f(r.profit)+" ₸"}</td><td class="r">${pc(r.margin)}</td><td class="r ${r.ret?"red":"mut"}">${r.ret?r.ret+" ("+pc(r.ret_pct)+")":"—"}</td>
    <td class="r">${r.shows?f(r.shows):"—"}</td><td class="r">${amb(r.ctr,3.5)}</td><td class="r">${amb(r.cr,2.5)}</td><td class="r">${r.ads?f(r.ads)+" ₸":"—"}</td>
    <td class="r">${r.drr==null?"—":`<span class="${r.drr>cap?"red":"grn"}">${pc(r.drr)}</span>`}</td><td class="r ${r.net_margin<0?"red":""}">${pc(r.net_margin)}</td></tr>`).join("")||'<tr><td colspan="15" class="msg">Ничего не найдено.</td></tr>';
    $("#pgb").innerHTML=`<span class="mut" style="padding:8px">Товаров: ${R.length}</span><button id="pv" ${PL.pg?"":"disabled"}>Назад</button><span class="mut" style="padding:8px">${PL.pg+1} из ${pages}</span><button id="nx" ${PL.pg+1<pages?"":"disabled"}>Вперёд</button>
    <select id="sz">${[10,25,50,100].map(n=>`<option ${n===PL.sz?"selected":""} value="${n}">${n} / стр.</option>`).join("")}</select>`;
    $("#pv").onclick=()=>{PL.pg--;draw()};$("#nx").onclick=()=>{PL.pg++;draw()};$("#sz").onchange=e=>{PL.sz=+e.target.value;PL.pg=0;draw()}};
  document.querySelectorAll("[data-pd]").forEach(x=>x.onclick=()=>{PR={from:ago(x.dataset.pd-1),to:iso(new Date())};pricelist()});
  $("#f").onchange=e=>{PR.from=e.target.value;pricelist()};$("#t").onchange=e=>{PR.to=e.target.value;pricelist()};$("#rf").onclick=()=>pricelist();
  document.querySelectorAll("[data-o]").forEach(x=>x.onclick=()=>{PL.d=PL.k===x.dataset.o?-PL.d:-1;PL.k=x.dataset.o;pricelist()});
  const reload=async()=>{d=(await api("/api/pricelist",PR)).data;draw()};
  $("#tb").onchange=async e=>{const t=e.target;await post("/api/cost",{sku:t.dataset.sku,field:t.dataset.k,value:t.value});reload()};
  document.querySelectorAll(".s").forEach(x=>x.onchange=async()=>{await post("/api/setting",{k:x.dataset.s,value:x.value});reload()});
  $("#s").oninput=e=>{PL.q=e.target.value;PL.pg=0;draw()};
  $("#ex").onclick=()=>{const q=v=>'"'+String(v??"").replace(/"/g,'""')+'"',keys=["sku","name","price","units","sold","revenue","cost","profit","margin","ret","shows","ctr","cr","ads","drr","net_margin"];
    const csv="\ufeff"+keys.join(";")+"\n"+view().map(r=>keys.map(k=>q(typeof r[k]=="number"?String(Math.round(r[k]*10)/10).replace(".",","):r[k])).join(";")).join("\n");
    const a=document.createElement("a");a.href=URL.createObjectURL(new Blob([csv],{type:"text/csv"}));a.download="prays-list-"+PR.from+"_"+PR.to+".csv";a.click()};
  draw()}
'''
MK_PANEL = r'''<div class="panel"><h2>Статистика рекламы по товару <span class="mut" style="font-weight:400">· из рекламного кабинета Kaspi, попадает в «Прайс-лист»</span></h2>
  <div class="bar" style="margin:0 0 8px"><input type="date" id="sd2" value="${iso(new Date())}"><input id="sk" placeholder="SKU товара"><input id="sh" type="number" placeholder="Показы"><input id="cl" type="number" placeholder="Клики"><input id="or" type="number" placeholder="Заказы"><input id="sp" type="number" step="any" placeholder="Расход, ₸"><button id="sb">Добавить</button></div>
  <div class="bar" style="margin:0"><input type="file" id="fl" accept=".xlsx,.csv"><button id="imp">Загрузить отчёт</button><span class="mut" id="impst" style="padding:8px">Колонки: дата, SKU или товар, показы, клики, заказы, расход</span></div></div>
  '''
MK_JS = r'''$("#sb").onclick=async()=>{const r=await post("/api/adstat",{day:$("#sd2").value,sku:$("#sk").value,spend:$("#sp").value,shows:$("#sh").value,clicks:$("#cl").value,orders:$("#or").value});if(r.error){alert(r.error);return}route()};
  $("#imp").onclick=()=>{const fl=$("#fl").files[0];if(!fl)return;const rd=new FileReader();rd.onload=async()=>{const r=await post("/api/ad/import",{name:fl.name,data:rd.result.split(",")[1]});$("#impst").textContent=r.error?"Ошибка: "+r.error:"Загружено строк: "+(r.n??"готово");if(!r.error)setTimeout(route,1200)};rd.readAsDataURL(fl)};
  '''
def need(c, msg):
    if not c: sys.exit("Не применено: " + msg)

s = open("server.py", encoding="utf-8").read(); shutil.copy("server.py", "server.py.bak3")
s = s.replace('{os.environ.get("PORT", 8000)}', "{os.environ.get('PORT', 8000)}")  # старый Python не принимает такие кавычки в f-строке
m = re.search(r"def pricelist\(.*?(?=\n\n\n)", s, re.S); need(m, "нет функции pricelist")
s = s[:m.start()] + SERVER_PRICELIST.rstrip("\n") + s[m.end():]
need('        elif path == "/api/syncnow":' in s, "нет ветки /api/syncnow"); s = s.replace('        elif path == "/api/syncnow":', POST_NEW, 1)
ALT = 'for _c in ("sku text", "shows integer", "clicks integer", "orders integer"):\n    try:\n        db.execute(f"alter table ads add column {_c}")\n    except sqlite3.OperationalError:\n        pass\ndb.commit()\n\n\n'
mm = re.search(r"\nif __name__ == .__main__.:", s); need(mm, "нет блока запуска if __name__")
s = s[:mm.start() + 1] + ALT + s[mm.start() + 1:]
s = re.sub(r'return \{"ok": True\}', 'return {"ok": True, "n": _n.get("n")}', s, count=1)
s = s.replace("if n > 100_000:", "if n > 5_000_000:")
try: ast.parse(s)
except SyntaxError as e: sys.exit(f"Ошибка синтаксиса в server.py, строка {e.lineno}: {e.msg}. Файл не изменён: пришлите мне эту строку.")
h = open("index.html", encoding="utf-8").read(); shutil.copy("index.html", "index.html.bak3")
m = re.search(r"let PR=.*?(?=\nasync function marketing\(\))", h, re.S) or re.search(r"async function pricelist\(\)\{.*?(?=\nasync function marketing\(\))", h, re.S)
need(m, "нет pricelist в index.html"); h = h[:m.start()] + JS_PRICE.rstrip("\n") + h[m.end():]
a = '<div class="panel tw"><h2>По дням</h2>'; need(a in h, "нет блока «По дням»"); h = h.replace(a, MK_PANEL + a, 1)
a = '$("#ab").onclick='; need(a in h, "нет кнопки добавления расхода"); h = h.replace(a, MK_JS + a, 1)
open("server.py", "w", encoding="utf-8").write(s); open("index.html", "w", encoding="utf-8").write(h)
print("Готово. Перезапустите сервер.")
