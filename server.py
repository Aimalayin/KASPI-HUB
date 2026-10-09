import os, re, io, json, time, sqlite3, threading, requests, hmac, base64, glob, sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

TOKEN = os.environ["KASPI_TOKEN"]
DIR = os.path.dirname(os.path.abspath(__file__))
DATA = os.environ.get("DATA_DIR", DIR)  # на хостинге: постоянный диск (например /data)
os.makedirs(DATA, exist_ok=True)
BACKFILL = int(os.environ.get("BACKFILL_DAYS", "30"))  # за сколько дней загрузить историю при запуске
wake = threading.Event()  # будит синхронизацию раньше срока
TZ = timezone(timedelta(hours=5))  # Алматы
BASE = "https://kaspi.kz/shop/api/v2"
H = {"X-Auth-Token": TOKEN, "Content-Type": "application/vnd.api+json", "User-Agent": "Mozilla/5.0"}
STATES = ["NEW", "SIGN_REQUIRED", "PICKUP", "DELIVERY", "KASPI_DELIVERY", "ARCHIVE"]
LIVE = "status not in ('CANCELLED','CANCELLING')"
REV = "status not in ('CANCELLED','CANCELLING','RETURNED','KASPI_DELIVERY_RETURN_REQUESTED','RETURN_ACCEPTED_BY_MERCHANT')"  # выручка без отмен и возвратов
PASSWORD = os.environ.get("HUB_PASSWORD", "")
USERS = [tuple(x.split(":", 1)) for x in os.environ.get("HUB_USERS", "").split(",") if ":" in x]  # «имя:пароль,имя2:пароль2» для разных людей
_fails = {}
BK = os.path.join(os.environ.get("DATA_DIR", os.path.dirname(os.path.abspath(__file__))), "backups")
RETS = "status in ('RETURNED','KASPI_DELIVERY_RETURN_REQUESTED','RETURN_ACCEPTED_BY_MERCHANT')"
LABEL = ("case when status in ('CANCELLED','CANCELLING') then 'Отменён' when " + RETS + " then 'Возврат' "
         "when status='COMPLETED' then 'Завершён' when pre=1 and state!='ARCHIVE' then 'Предзаказ' "
         "when state in ('NEW','SIGN_REQUIRED') then 'Новый' when state in ('DELIVERY','KASPI_DELIVERY') then 'В доставке' "
         "when state='PICKUP' then 'Самовывоз' else 'Архив' end")

db = sqlite3.connect(os.path.join(DATA, "hub.db"), check_same_thread=False)
db.execute("pragma journal_mode=wal")
lock = threading.Lock()
db.executescript("""
create table if not exists orders(id text primary key, code text, state text, status text, total real,
  created integer, city text, customer text, phone text, pre integer, items integer default 0);
create table if not exists entries(oid text, name text, code text, qty integer, total real);
create index if not exists i_created on orders(created);
create index if not exists i_oid on entries(oid);
create table if not exists costs(sku text primary key, cost real, delivery real, mkt real, comm real);
create table if not exists settings(k text primary key, v real);
create table if not exists ads(id integer primary key autoincrement, day text, name text, spend real);
create table if not exists prefs(k text primary key, v text);
create table if not exists log(id integer primary key autoincrement, t integer, kind text, text text);""")
for col in ("planned", "assembled"):
    try:
        db.execute(f"alter table orders add column {col} integer")
    except sqlite3.OperationalError:
        pass
info = {"sync": "первая загрузка…", "error": None}


def q(sql, a=()):
    with lock:
        return db.execute(sql, a).fetchall()


def api(path, params):
    for _ in range(3):
        try:
            r = requests.get(BASE + path, headers=H, params=params, timeout=120)
            if r.status_code == 200:
                return r.json()
            if r.status_code == 429:
                time.sleep(10)
            if r.status_code in (401, 403):
                raise PermissionError(f"Kaspi вернул {r.status_code}: проверьте токен")
        except requests.RequestException:
            pass
        time.sleep(2)
    raise ConnectionError("Kaspi не отвечает")


def ms(d):
    return int(d.timestamp() * 1000)


def sync_orders(days):
    day0 = datetime.now(TZ).replace(hour=0, minute=0, second=0, microsecond=0)
    stop, s = day0 + timedelta(days=1), day0 - timedelta(days=days - 1)
    while s < stop:
        e = min(s + timedelta(days=7), stop)
        for st in STATES:
            p = 0
            while True:
                d = api("/orders", {"page[number]": p, "page[size]": 100, "filter[orders][state]": st,
                                    "filter[orders][creationDate][$ge]": ms(s), "filter[orders][creationDate][$le]": ms(e) - 1})
                rows = []
                for o in d["data"]:
                    a, c = o["attributes"], o["attributes"].get("customer") or {}
                    nm = c.get("name") or f"{c.get('firstName') or ''} {c.get('lastName') or ''}".strip()
                    rows.append((o["id"], a.get("code"), a.get("state"), a.get("status"), a.get("totalPrice", 0),
                                 a["creationDate"], (a.get("deliveryAddress") or {}).get("town"), nm,
                                 c.get("cellPhone"), 1 if a.get("preOrder") else 0, a.get("plannedDeliveryDate"), 1 if a.get("assembled") else 0))
                with lock:
                    db.executemany("insert into orders(id,code,state,status,total,created,city,customer,phone,pre,planned,assembled) "
                                   "values(?,?,?,?,?,?,?,?,?,?,?,?) on conflict(id) do update set "
                                   "state=excluded.state,status=excluded.status,total=excluded.total,planned=excluded.planned,assembled=excluded.assembled", rows)
                    db.commit()
                p += 1
                if p >= d["meta"]["pageCount"]:
                    break
        s = e


def sync_loop():
    last = 0
    while True:
        try:
            full = time.time() - last > 3600  # раз в час обновляем статусы за весь период
            days = full_days() if full else 3
            sync_orders(days)
            if full:
                last = time.time()
                log("Синхронизация", f"Заказы обновлены за {days} дн.")
            info.update(sync=datetime.now(TZ).strftime("%H:%M"), error=None)
            print("заказы обновлены", info["sync"], flush=True)
        except Exception as e:
            if info["error"] != str(e):
                log("Синхронизация", f"Ошибка: {e}")
            info["error"] = str(e)
            print("ошибка:", e, flush=True)
        if wake.wait(max(2, int(pf("refresh_min"))) * 60):
            wake.clear()
            last = 0  # следующий проход полный


_seen = {"sample": False, "err": 0}


def fetch_entries(oid):
    try:
        d = api(f"/orders/{oid}/entries", {"include[orderentries]": "product,merchantProduct"})
    except PermissionError:
        raise
    except Exception as e:
        if _seen["err"] < 3:
            _seen["err"] += 1
            log("Синхронизация", f"Не загрузились позиции заказа {oid}: {e}")
        return
    inc = {(x.get("type"), x["id"]): x.get("attributes") or {} for x in d.get("included", [])}

    def pick(e, k):
        r = ((e.get("relationships") or {}).get(k) or {}).get("data") or {}
        return inc.get((r.get("type"), r.get("id"))) or next((v for (t, i), v in inc.items() if i == r.get("id")), {})
    rows = []
    for e in d["data"]:
        a = e["attributes"]
        mp, p = pick(e, "merchantProduct"), pick(e, "product")
        code = mp.get("code") or (a.get("offer") or {}).get("code") or p.get("code")
        name = mp.get("name") or p.get("name") or (a.get("category") or {}).get("title") or "Без названия"
        if not code and not _seen["sample"]:  # диагностика: сырой ответ попадёт в «Журнал действий»
            _seen["sample"] = True
            log("Диагностика", "Kaspi не отдал SKU позиции. Сырой ответ: " + json.dumps(d, ensure_ascii=False)[:1500])
        rows.append((oid, name, str(code) if code else None, a.get("quantity", 1), a.get("totalPrice", 0)))
    with lock:
        db.execute("delete from entries where oid=?", (oid,))
        db.executemany("insert into entries values(?,?,?,?,?)", rows)
        db.execute("update orders set items=1 where id=?", (oid,))
        db.commit()


def entries_loop():
    while True:
        ids = [r[0] for r in q(f"select id from orders where items=0 and {LIVE} order by created desc limit 120")]
        if not ids:
            time.sleep(10)
            continue
        try:
            with ThreadPoolExecutor(6) as ex:
                list(ex.map(fetch_entries, ids))
        except PermissionError as e:
            info["error"] = str(e)
            time.sleep(60)


_stock = {"mtime": 0, "rows": []}


def load_stock():
    p = os.path.join(DATA, "ACTIVE.xlsx")
    if not os.path.exists(p):
        p = os.path.join(DIR, "ACTIVE.xlsx")
    if not os.path.exists(p):
        return []
    if os.path.getmtime(p) != _stock["mtime"]:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True)
        rows = []
        for r in list(wb.worksheets[0].iter_rows(values_only=True))[1:]:
            if not r or not r[0]:
                continue
            units = 0
            for x in r[4:9]:
                try:
                    units += max(0, int(float(str(x).strip().replace(",", "."))))
                except ValueError:
                    pass
            try:
                price, pre = float(r[3] or 0), int(float(r[9] or 0))
            except (ValueError, IndexError):
                price, pre = 0.0, 0
            rows.append({"sku": str(r[0]).strip(), "name": r[1] or "—", "brand": r[2] or "", "price": price, "units": units, "pre": pre})
        wb.close()
        _stock.update(mtime=os.path.getmtime(p), rows=rows)
    return _stock["rows"]


def match(code, name, S):
    """Позиция заказа -> товар склада: по SKU (в т.ч. по части до «_»), затем по уникальному названию."""
    bc, bp, bn = {}, {}, {}
    for s in S:
        bc[s["sku"]] = s
        bp.setdefault(s["sku"].split("_")[0], []).append(s)
        bn.setdefault(s["name"].lower(), []).append(s)
    if code:
        if code in bc:
            return bc[code]
        c = bp.get(code.split("_")[0], [])
        if len(c) == 1:
            return c[0]
    c = bn.get((name or "").lower(), [])
    return c[0] if len(c) == 1 else None


def sold(a, b):
    # продано без возвратов; последний столбец: сколько возвращено
    return q(f"select e.name,e.code,sum(case when o.{RETS} then 0 else e.qty end),sum(case when o.{RETS} then 0 else e.total end),"
             f"sum(case when o.{RETS} then e.qty else 0 end) from entries e join orders o on o.id=e.oid "
             f"where o.created between ? and ? and o.{LIVE} group by e.name,e.code", (a, b))


_rc = {"t": 0, "v": {}}


def rates():
    if time.time() - _rc["t"] < 60:
        return _rc["v"]
    now = ms(datetime.now(TZ))
    S, out = load_stock(), {}
    for name, code, qty, _, _ in sold(now - 30 * 864e5, now):
        s = match(code, name, S)
        if s:
            out[s["sku"]] = out.get(s["sku"], 0) + qty / 30
    _rc.update(t=time.time(), v=out)
    return out


def rng(p):
    t = str(datetime.now(TZ).date())
    a = datetime.fromisoformat(p.get("from", [t])[0]).replace(tzinfo=TZ)
    b = datetime.fromisoformat(p.get("to", [t])[0]).replace(tzinfo=TZ) + timedelta(days=1)
    return ms(a), ms(b) - 1


def pending(a, b):
    return q(f"select count(*) from orders where items=0 and {LIVE} and created between ? and ?", (a, b))[0][0]


def totals(a, b):
    n, rev = q(f"select count(*), coalesce(sum(case when {REV} then total end),0) from orders where created between ? and ?", (a, b))[0]
    live = q(f"select count(*) from orders where {REV} and created between ? and ?", (a, b))[0][0]
    labels = {l: c for l, c in q(f"select {LABEL} l, count(*) from orders where created between ? and ? group by l", (a, b))}
    return {"total": n, "revenue": rev, "avg": round(rev / live) if live else 0, "labels": labels}


def today():
    a, b = rng({})
    t = totals(a, b)
    hours = [[0, 0] for _ in range(24)]
    for h, c, r in q(f"select strftime('%H', created/1000+18000, 'unixepoch'), count(*), coalesce(sum(case when {REV} then total end),0) from orders where created between ? and ? group by 1", (a, b)):
        hours[int(h)] = [c, r]
    S, agg = load_stock(), {}
    for name, code, qty, total, _ in sold(a, b):  # группируем по товару склада, а не по категории Kaspi
        sk = match(code, name, S)
        r = agg.setdefault(sk["sku"] if sk else (code or name), {"name": sk["name"] if sk else name, "qty": 0, "sum": 0})
        r["qty"] += qty
        r["sum"] += total
    top = sorted(agg.values(), key=lambda r: (-r["qty"], -r["sum"]))  # все товары за сегодня, фронт показывает 5 + «Больше»
    s = stock()
    return {**t, "hours": hours, "top": top, "pending": pending(a, b), "health": health(s), "recs": recs(s)}


def analytics(p):
    a, b = rng(p)
    t = totals(a, b)
    daily = [{"d": d, "rev": r, "n": n} for d, r, n in q(
        f"select strftime('%Y-%m-%d', created/1000+18000, 'unixepoch') d, coalesce(sum(case when {REV} then total end),0), count(*) "
        "from orders where created between ? and ? group by d order by d", (a, b))]
    rt = rates()
    att, dd = [], pf("deficit_days")
    for s in load_stock():
        r = rt.get(s["sku"], 0)
        if r > 0 and s["units"] == 0:
            att.append(f"Нет в наличии: {s['name']}: продаётся {r:.1f} шт/день")
        elif r > 0 and s["units"] / r < dd:
            att.append(f"Скоро закончится: {s['name']}: хватит на {s['units'] / r:.0f} дн")
    return {**t, "daily": daily, "attention": att[:10], "cancelled": t["labels"].get("Отменён", 0),
            "returns": t["labels"].get("Возврат", 0)}


def products(p):
    a, b = rng(p)
    S, agg = load_stock(), {}
    for name, code, qty, total, ret in sold(a, b):
        s = match(code, name, S)
        k = s["sku"] if s else (code or name)
        r = agg.setdefault(k, {"name": s["name"] if s else name, "sku": s["sku"] if s else (code or ""), "price": s and s["price"],
                               "units": s and s["units"], "est": False, "sold": 0, "revenue": 0, "returns": 0})
        r["sold"] += qty
        r["revenue"] += total
        r["returns"] += ret
    for r in agg.values():  # товара нет в прайс-листе (закончился/снят): берём среднюю цену продажи из заказов
        if r["price"] is None and r["sold"]:
            r["price"] = round(r["revenue"] / r["sold"])
            r["est"] = True
    rows = list(agg.values())
    rows += [{"name": s["name"], "sku": s["sku"], "price": s["price"], "units": s["units"], "est": False, "sold": 0, "revenue": 0, "returns": 0}
             for s in S if s["sku"] not in agg]
    rows.sort(key=lambda r: -r["revenue"])
    return {"rows": rows, "sold": sum(r["sold"] for r in rows), "revenue": sum(r["revenue"] for r in rows), "pending": pending(a, b)}


def orders(p):
    a, b = rng(p)
    like, lab = p.get("q", [""])[0], p.get("label", [""])[0]
    page = int(p.get("page", ["0"])[0])
    w, args = "created between ? and ?", [a, b]
    if like:
        w += " and (code like ? or customer like ? or phone like ? or city like ?)"
        args += [f"%{like}%"] * 4
    if lab:
        w += f" and {LABEL}=?"
        args.append(lab)
    n = q(f"select count(*) from orders where {w}", args)[0][0]
    rows = q(f"select code,created,{LABEL},customer,phone,city,total,(select coalesce(sum(qty),0) from entries where oid=orders.id) "
             f"from orders where {w} order by created desc limit 50 offset ?", args + [page * 50])
    return {"n": n, "rows": [{"code": r[0], "t": datetime.fromtimestamp(r[1] / 1000, TZ).strftime("%d.%m %H:%M"), "label": r[2],
                              "customer": r[3], "phone": None if (r[4] or "").startswith("+0(000)") else r[4], "city": r[5], "total": r[6], "items": r[7]} for r in rows]}


def stock():
    S, rt, dd = load_stock(), rates(), pf("deficit_days")
    rows = []
    for s in S:
        r = rt.get(s["sku"], 0)
        days = s["units"] / r if r > 0 else None
        rows.append({**s, "rate": round(r, 2), "days": None if days is None else round(days, 1),
                     "value": s["units"] * s["price"], "out": s["units"] == 0,
                     "deficit": s["units"] == 0 and r > 0 or (days is not None and days < dd)})
    return {"sku": len(S), "units": sum(r["units"] for r in rows), "value": sum(r["value"] for r in rows),
            "in": sum(1 for r in rows if not r["out"]), "out": sum(1 for r in rows if r["out"]),
            "deficit": sum(1 for r in rows if r["deficit"]), "rows": rows}


def health(s):
    bad = [r for r in s["rows"] if r["deficit"]]
    out = [r for r in s["rows"] if r["out"] and not r["deficit"]]
    items = sorted(bad + out, key=lambda r: r["units"])[:8]
    return {"ok": len(s["rows"]) - len(bad) - len(out), "deficit": len(bad), "out": len(out), "value": s["value"],
            "items": [{"sku": r["sku"], "name": r["name"], "units": r["units"]} for r in items]}


def recs(s):
    now, out = ms(datetime.now(TZ)), []
    bad = [r for r in s["rows"] if r["deficit"]]
    if bad:
        out.append({"level": "Критично", "title": f"Дефицит склада · {len(bad)}",
                    "text": " · ".join(f"{r['name']} ({r['units']})" for r in bad[:3])})
    n = q(f"select count(*) from orders where pre=1 and planned<? and state!='ARCHIVE' and {LIVE} and status!='COMPLETED'", (now,))[0][0]
    if n:
        out.append({"level": "Важно", "title": f"Просрочено предзаказов · {n}", "text": "Плановая дата прошла: продлите или отмените"})
    for code, days in q(f"select code,(?-created)/86400000 from orders where state in ('DELIVERY','KASPI_DELIVERY') and {LIVE} "
                        "and status!='COMPLETED' and created<? order by created limit 3", (now, now - pf("stuck_days") * 864e5)):
        out.append({"level": "Важно", "title": f"Долго в доставке · заказ {code}",
                    "text": f"В доставке уже {int(days)} дн. Проверьте статус и при необходимости подайте претензию в Kaspi"})
    n2 = q("select count(*) from orders where status='CANCELLING' and created<?", (now - 3 * 864e5,))[0][0]
    if n2:
        out.append({"level": "Важно", "title": f"Зависших отмен · {n2}", "text": "Заказы в статусе «Отменяется» больше 3 дней: проверьте в кабинете Kaspi"})
    cap = pf("drr_cap")
    if cap:
        a = now - 7 * 864e5
        day = lambda x: datetime.fromtimestamp(x / 1000, TZ).strftime("%Y-%m-%d")
        sp = q("select coalesce(sum(spend),0) from ads where day between ? and ?", (day(a), day(now)))[0][0]
        rev = totals(a, now)["revenue"]
        if sp and rev and sp / rev * 100 > cap:
            out.append({"level": "Важно", "title": f"ДРР выше потолка · {sp / rev * 100:.1f}%", "text": f"За 7 дней реклама съела больше {cap:g}% выручки"})
    return out


PREF_DEF = {"company": "Kaspi Hub", "start": "", "refresh_min": "5", "deficit_days": "7", "stuck_days": "9", "drr_cap": "0"}
NUM_PREFS = ("refresh_min", "deficit_days", "stuck_days", "drr_cap")
PLBL = {"company": "Название компании", "start": "Начало синхронизации заказов", "refresh_min": "Интервал обновления, мин",
        "deficit_days": "Порог дефицита, дней", "stuck_days": "Долго в доставке, дней", "drr_cap": "Потолок ДРР, %"}
LBL = {"cost": "себестоимость", "delivery": "доставка ₸", "mkt": "маркетинг %", "comm": "комиссия %"}
SLBL = {"comm": "комиссия Kaspi %", "delivery": "доставка ₸", "mkt": "маркетинг %", "tax": "налог %"}


def fmtv(v):
    return "—" if v is None else f"{v:g}"


def _log(kind, text):  # вызывать, когда lock уже взят
    db.execute("insert into log(t,kind,text) values(?,?,?)", (ms(datetime.now(TZ)), kind, text))


def log(kind, text):
    with lock:
        _log(kind, text)
        db.commit()


_pc = {"t": 0, "v": None}


def prefs():
    if _pc["v"] is None or time.time() - _pc["t"] > 5:
        g = dict(PREF_DEF)
        g.update(dict(q("select k,v from prefs")))
        _pc.update(t=time.time(), v=g)
    return dict(_pc["v"])


def pf(k):
    try:
        return float(prefs()[k])
    except ValueError:
        return float(PREF_DEF[k])


def full_days():
    try:
        return max(1, (datetime.now(TZ).date() - datetime.fromisoformat(prefs()["start"]).date()).days + 1)
    except ValueError:
        return BACKFILL


def logs(p):
    kind, like, page = p.get("kind", [""])[0], p.get("q", [""])[0], int(p.get("page", ["0"])[0])
    w, args = "1=1", []
    if kind:
        w += " and kind=?"
        args.append(kind)
    if like:
        w += " and text like ?"
        args.append(f"%{like}%")
    rows = q(f"select t,kind,text from log where {w} order by id desc limit 50 offset ?", args + [page * 50])
    return {"n": q(f"select count(*) from log where {w}", args)[0][0], "kinds": [r[0] for r in q("select distinct kind from log order by 1")],
            "rows": [{"t": datetime.fromtimestamp(r[0] / 1000, TZ).strftime("%d.%m %H:%M:%S"), "kind": r[1], "text": r[2]} for r in rows]}


def getset():
    g = {"comm": 0, "delivery": 0, "mkt": 0, "tax": 0}
    g.update({k: v for k, v in q("select k,v from settings")})
    return g


def pricelist(p):
    a, b = rng(p)
    g, S = getset(), load_stock()
    C = {r[0]: r for r in q("select sku,cost,delivery,mkt,comm from costs")}
    z = lambda: {"sold": 0, "revenue": 0, "ret": 0, "spend": 0, "shows": 0, "clicks": 0, "aord": 0}
    agg = {x["sku"]: z() for x in S}
    for name, code, qty, total, ret in sold(a, b):
        x = match(code, name, S)
        if x:
            r = agg[x["sku"]]
            r["sold"] += qty
            r["revenue"] += total
            r["ret"] += ret
    day = lambda t: datetime.fromtimestamp(t / 1000, TZ).strftime("%Y-%m-%d")
    tot = z()
    for nm, sk, sp, sh, cl, od in q("select name,sku,spend,shows,clicks,orders from ads where day between ? and ?", (day(a), day(b))):
        v = {"spend": sp or 0, "shows": sh or 0, "clicks": cl or 0, "aord": od or 0}
        for k in v:
            tot[k] += v[k]
        key = None
        if sk:
            x = match(str(sk), "", S)
            key = x["sku"] if x else None
        if key is None:  # реклама привязывается к товару по SKU или если в названии есть SKU/название товара
            n = (nm or "").lower()
            for x in S:
                if x["sku"].lower() in n or x["name"].lower() in n:
                    key = x["sku"]
                    break
        if key:
            for k in v:
                agg[key][k] += v[k]
    rows = []
    for x in S:
        _, cost, od, om, oc = C.get(x["sku"], (None,) * 5)
        v = agg[x["sku"]]
        cm, dl = (g["comm"] if oc is None else oc), (g["delivery"] if od is None else od)
        rev, n = v["revenue"], v["sold"]
        profit = None if cost is None else rev * (1 - (cm + g["tax"]) / 100) - n * (cost + dl)
        ad = v["spend"]
        net = None if profit is None else profit - ad
        rows.append({"sku": x["sku"], "name": x["name"], "price": x["price"], "units": x["units"], "sold": n, "revenue": rev,
                     "cost": cost, "cost_total": None if cost is None else n * cost, "profit": profit,
                     "margin": profit / rev * 100 if profit is not None and rev else None,
                     "ret": v["ret"], "ret_pct": v["ret"] / (n + v["ret"]) * 100 if n + v["ret"] else None,
                     "shows": v["shows"] or None, "ctr": v["clicks"] / v["shows"] * 100 if v["shows"] else None,
                     "cr": v["aord"] / v["clicks"] * 100 if v["clicks"] and v["aord"] else None,
                     "ads": ad, "drr": ad / rev * 100 if ad and rev else None, "net": net,
                     "net_margin": net / rev * 100 if net is not None and rev else None})
    kn = [r for r in rows if r["profit"] is not None]
    rk, pr = sum(r["revenue"] for r in kn), sum(r["profit"] for r in kn)
    k = {"sold": sum(r["sold"] for r in rows), "revenue": sum(r["revenue"] for r in rows), "cost": sum(r["cost_total"] for r in kn),
         "profit": pr, "margin": pr / rk * 100 if rk else None, "ads": tot["spend"], "net": pr - tot["spend"],
         "net_margin": (pr - tot["spend"]) / rk * 100 if rk else None,
         "ctr": tot["clicks"] / tot["shows"] * 100 if tot["shows"] else None,
         "cr": tot["aord"] / tot["clicks"] * 100 if tot["clicks"] and tot["aord"] else None,
         "no_cost": sum(1 for r in rows if r["cost"] is None and r["sold"])}
    return {"settings": g, "rows": rows, "k": k}


_n = {}


def import_ads(fname, b64):
    """Отчёт из рекламного кабинета (xlsx/csv): колонки ищутся по названию: дата, SKU/товар, показы, клики, заказы, расход."""
    import io, csv
    raw = base64.b64decode(b64)
    try:
        if fname.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
            wb.close()
        else:
            txt = raw.decode("utf-8-sig", errors="replace")
            rows = list(csv.reader(io.StringIO(txt), delimiter=";" if txt.count(";") > txt.count(",") else ","))
    except Exception as e:
        raise ValueError(f"Не удалось прочитать файл: {e}")
    syn = {"day": ("дата", "date", "день"), "sku": ("sku", "артикул", "код"), "name": ("товар", "название", "наименование", "кампания", "модель", "product", "name"),
           "shows": ("показ", "просмотр", "impression", "shows"), "clicks": ("клик", "переход", "click"), "orders": ("заказ", "order", "покупк"),
           "spend": ("расход", "затрат", "потрачен", "стоимост", "бюджет", "spend", "cost", "сумма")}
    hdr = None
    for i, r in enumerate(rows[:15]):
        cells = [str(c or "").strip().lower() for c in r]
        m = {k: next((j for j, c in enumerate(cells) if any(c.startswith(s) for s in v)), None) for k, v in syn.items()}
        if m["spend"] is not None or m["shows"] is not None:
            hdr = (i, m)
            break
    if not hdr:
        raise ValueError("Не нашёл колонки: нужны хотя бы «Показы» или «Расход»")
    i0, m = hdr

    def num(x):
        try:
            return float(str(x).replace("\xa0", "").replace(" ", "").replace(",", ".")) if x not in (None, "") else 0.0
        except ValueError:
            return 0.0

    def dt(x):
        if hasattr(x, "strftime"):
            return x.strftime("%Y-%m-%d")
        s = str(x or "").strip()[:10]
        for f in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(s, f).strftime("%Y-%m-%d")
            except ValueError:
                pass
        return datetime.now(TZ).strftime("%Y-%m-%d")
    cnt = 0
    for r in rows[i0 + 1:]:
        g = lambda k: r[m[k]] if m[k] is not None and m[k] < len(r) else None
        sku, nm = str(g("sku") or "").strip(), str(g("name") or "").strip()
        if not (sku or nm) or str(sku + nm).lower().startswith(("итого", "total")):
            continue
        d = dt(g("day"))
        db.execute("delete from ads where day=? and coalesce(sku,name)=?", (d, sku or nm[:80]))
        db.execute("insert into ads(day,name,spend,sku,shows,clicks,orders) values(?,?,?,?,?,?,?)",
                   (d, (nm or sku)[:80], num(g("spend")), sku or None, int(num(g("shows"))), int(num(g("clicks"))), int(num(g("orders")))))
        cnt += 1
    if not cnt:
        raise ValueError("В файле нет строк с товарами")
    return cnt


def marketing(p):
    a, b = rng(p)
    day = lambda x: datetime.fromtimestamp(x / 1000, TZ).strftime("%Y-%m-%d")
    ads = [{"id": r[0], "day": r[1], "name": r[2], "spend": r[3]} for r in
           q("select id,day,name,spend from ads where day between ? and ? order by day desc, id desc", (day(a), day(b)))]
    t = totals(a, b)
    days = {}
    for d, rev, n in q(f"select strftime('%Y-%m-%d', created/1000+18000, 'unixepoch') d, coalesce(sum(case when {REV} then total end),0), "
                       f"sum(case when {LIVE} then 1 else 0 end) from orders where created between ? and ? group by d", (a, b)):
        days[d] = {"day": d, "rev": rev, "n": n, "spend": 0}
    for x in ads:
        days.setdefault(x["day"], {"day": x["day"], "rev": 0, "n": 0, "spend": 0})["spend"] += x["spend"]
    return {"ads": ads, "spend": sum(x["spend"] for x in ads), "revenue": t["revenue"],
            "orders": t["total"] - t["labels"].get("Отменён", 0), "cap": pf("drr_cap"), "daily": sorted(days.values(), key=lambda r: r["day"], reverse=True)}


def upload_stock(d):
    import io
    from openpyxl import load_workbook
    raw = base64.b64decode(d["file"])
    assert len(raw) < 4_000_000, "Файл слишком большой"
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True)
        rows = [r for r in wb.worksheets[0].iter_rows(values_only=True) if r and r[0]]
    except Exception:
        raise ValueError("Это не файл .xlsx")
    if len(rows) < 2:
        raise ValueError("В файле нет таблицы товаров (ACTIVE.xlsx)")
    tmp = os.path.join(DATA, "ACTIVE.xlsx.tmp")
    open(tmp, "wb").write(raw)
    os.replace(tmp, os.path.join(DATA, "ACTIVE.xlsx"))
    _stock["mtime"] = None
    _rc["t"] = 0
    log("Склад", f"Загружен новый ACTIVE.xlsx: {len(rows) - 1} строк")
    return {"ok": True, "rows": len(rows) - 1}


def post(path, d):
    if path == "/api/stock_upload":
        return upload_stock(d)
    with lock:
        if path == "/api/cost":
            f = d["field"]
            assert f in ("cost", "delivery", "mkt", "comm")
            v = None if d.get("value") in ("", None) else float(d["value"])
            old = db.execute(f"select {f} from costs where sku=?", (d["sku"],)).fetchone()
            nm = next((x["name"] for x in load_stock() if x["sku"] == d["sku"]), d["sku"])
            _log("Прайс-лист", f"{nm}: {LBL[f]} {fmtv(old[0] if old else None)} → {fmtv(v)}")
            db.execute("insert or ignore into costs(sku) values(?)", (d["sku"],))
            db.execute(f"update costs set {f}=? where sku=?", (v, d["sku"]))
        elif path == "/api/setting":
            assert d["k"] in ("comm", "delivery", "mkt", "tax")
            db.execute("insert or replace into settings values(?,?)", (d["k"], float(d["value"] or 0)))
            _log("Прайс-лист", f"Общая ставка «{SLBL[d['k']]}»: {float(d['value'] or 0):g}")
        elif path == "/api/ad":
            datetime.fromisoformat(d["day"])
            db.execute("insert into ads(day,name,spend) values(?,?,?)", (d["day"], str(d.get("name") or "Реклама")[:80], float(d["spend"])))
            _log("Маркетинг", f"Добавлен расход {float(d['spend']):g} ₸: {d.get('name') or 'Реклама'}, {d['day']}")
        elif path == "/api/ad/delete":
            r = db.execute("select day,name,spend from ads where id=?", (int(d["id"]),)).fetchone()
            if r:
                _log("Маркетинг", f"Удалён расход {r[2]:g} ₸: {r[1]}, {r[0]}")
            db.execute("delete from ads where id=?", (int(d["id"]),))
        elif path == "/api/prefs":
            cur = dict(PREF_DEF)
            cur.update(dict(db.execute("select k,v from prefs").fetchall()))
            for k, v in d.items():
                assert k in PREF_DEF
                v = str(v).strip()
                if k == "start" and v:
                    datetime.fromisoformat(v)
                if k in NUM_PREFS:
                    v = f"{float(v):g}" if v else PREF_DEF[k]
                if v != cur[k]:
                    _log("Настройки", f"{PLBL[k]}: {cur[k] or '—'} → {v or '—'}")
                    db.execute("insert or replace into prefs values(?,?)", (k, v))
                    if k == "start":
                        wake.set()
        elif path == "/api/waybills":
            if job["state"] != "run":
                job.update(state="run", done=0, total=0, file=None, errors=[])
                threading.Thread(target=run_waybills, daemon=True).start()
        elif path == "/api/adstat":
            datetime.fromisoformat(d["day"])
            sk = str(d.get("sku") or "").strip()
            nm = str(d.get("name") or sk or "Реклама")[:80]
            db.execute("insert into ads(day,name,spend,sku,shows,clicks,orders) values(?,?,?,?,?,?,?)",
                       (d["day"], nm, float(d.get("spend") or 0), sk or None, int(float(d.get("shows") or 0)), int(float(d.get("clicks") or 0)), int(float(d.get("orders") or 0))))
            _log("Маркетинг", f"Статистика рекламы {nm}, {d['day']}: показы {d.get('shows') or 0}, клики {d.get('clicks') or 0}, расход {d.get('spend') or 0} ₸")
        elif path == "/api/ad/import":
            _n["n"] = import_ads(str(d["name"]), d["data"])
            _log("Маркетинг", f"Загружен отчёт по рекламе «{d['name']}»: строк {_n['n']}")
        elif path == "/api/syncnow":
            _log("Синхронизация", "Запущена вручную")
            wake.set()
        else:
            raise ValueError("неизвестный запрос")
        db.commit()
    _pc["t"] = 0
    _rc["t"] = 0
    return {"ok": True, "n": _n.get("n")}



# ---------- Упаковка: накладные одним PDF, одинаковые товары подряд ----------
WB = os.path.join(DIR, "waybills")
PACK = "state='KASPI_DELIVERY' and status='ACCEPTED_BY_MERCHANT' and (assembled=0 or id in (select oid from packed))"
job = {"state": "idle", "done": 0, "total": 0, "file": None, "errors": []}


def pack_sorted():
    S, out = load_stock(), []
    for oid, code, created in q(f"select id,code,created from orders where {PACK} order by created"):
        items = {}
        for name, c, qty in q("select name,code,qty from entries where oid=?", (oid,)):
            s = match(c, name, S)
            k = s["name"] if s else name
            items[k] = items.get(k, 0) + qty
        out.append({"id": oid, "code": code, "created": created, "items": items})
    units, cnt = {}, {}
    for o in out:
        if len(o["items"]) == 1:
            k = next(iter(o["items"]))
            units[k] = units.get(k, 0) + o["items"][k]
            cnt[k] = cnt.get(k, 0) + 1
    # сначала заказы с одним товаром: группы по убыванию штук, внутри группы подряд; смешанные и без позиций в конце
    out.sort(key=lambda o: (0, -units[next(iter(o["items"]))], next(iter(o["items"])), o["created"]) if len(o["items"]) == 1
             else (1, 0, " + ".join(sorted(o["items"])), o["created"]))
    groups = [{"name": k, "units": units[k], "orders": cnt[k]} for k in sorted(units, key=lambda k: (-units[k], k))]
    return out, groups


def pack_view():
    out, groups = pack_sorted()
    return {"orders": [{"code": o["code"], "items": ", ".join(f"{k} ×{v}" for k, v in o["items"].items()) or "позиции не загружены"} for o in out],
            "groups": groups, "mixed": sum(1 for o in out if len(o["items"]) != 1), "job": job}


def waybill_one(oid):
    code = (q("select code from orders where id=?", (oid,)) or [[None]])[0][0]

    def find(x):  # ищем именно ссылку (http или путь), а не номер накладной
        if isinstance(x, dict):
            for k, v in x.items():
                if isinstance(v, str) and v.startswith(("http", "/")) and ("waybill" in str(k).lower() or v.lower().split("?")[0].endswith(".pdf")):
                    return v
            for v in x.values():
                r = find(v)
                if r:
                    return r
        elif isinstance(x, list):
            for v in x:
                r = find(v)
                if r:
                    return r
        return None

    def kdel(got):  # что Kaspi отдал про доставку: для сообщения об ошибке
        for d in got:
            dd = d.get("data")
            dd = dd[0] if isinstance(dd, list) and dd else dd
            a = (dd or {}).get("attributes") or {} if isinstance(dd, dict) else {}
            if a.get("kaspiDelivery"):
                return a["kaspiDelivery"]
        return None

    def link():
        got = []
        for path, prm in ((f"/orders/{oid}", {}), ("/orders", {"page[number]": 0, "page[size]": 1, "filter[orders][code]": code})):
            if path == "/orders" and not code:
                continue
            try:
                d = api(path, prm)
            except PermissionError:
                raise
            except Exception:
                continue
            u = find(d)
            if u:
                return u, got
            got.append(d)
        return None, got
    url, got = link()
    if not url:  # то же, что кнопка «Сформировать» в Упаковке: после сборки Kaspi создаёт накладную
        body = {"data": {"type": "orders", "id": oid, "attributes": {"status": "ASSEMBLE", "numberOfSpace": str(max(1, int(os.environ.get("PACK_SPACES", "1"))))}}}
        r = requests.post(BASE + "/orders", headers=H, json=body, timeout=60)
        if r.status_code not in (200, 201, 204):
            raise ValueError(f"Kaspi не принял сборку (код {r.status_code}): {r.text[:200]}")
        with lock:  # заказ собран: помним, чтобы он не пропал из списка, пока накладная не скачана
            db.execute("create table if not exists packed(oid text primary key)")
            db.execute("insert or ignore into packed values(?)", (oid,))
            db.commit()
        try:
            url = find(r.json())
        except ValueError:
            pass
        for _ in range(6):
            if url:
                break
            time.sleep(3)
            url, got = link()
    if not url:
        if not _seen.get("wb"):
            _seen["wb"] = True
            log("Диагностика", "Нет ссылки на накладную. Ответ Kaspi: " + json.dumps(got, ensure_ascii=False)[:1800])
        raise ValueError("Kaspi не отдал ссылку на накладную. Про доставку Kaspi отдал: " + json.dumps(kdel(got), ensure_ascii=False)[:300])
    if url.startswith("/"):
        url = "https://kaspi.kz" + url
    host = urlparse(url).hostname or ""
    if not (host == "kaspi.kz" or host.endswith(".kaspi.kz")):
        raise ValueError("ссылка ведёт не на Kaspi: " + url[:60])
    r = requests.get(url, headers={"X-Auth-Token": TOKEN, "User-Agent": "Mozilla/5.0"}, timeout=60)
    if r.status_code != 200 or not r.content.startswith(b"%PDF"):
        raise ValueError(f"накладная не скачалась (код {r.status_code})")
    return r.content


def run_waybills():
    try:
        from pypdf import PdfWriter, PdfReader
        os.makedirs(WB, exist_ok=True)
        for (oid,) in q(f"select id from orders where {PACK} and not exists(select 1 from entries where oid=orders.id)"):
            fetch_entries(oid)  # докачиваем позиции, чтобы правильно сгруппировать
        orders, _ = pack_sorted()
        lim = int(os.environ.get("PACK_LIMIT", "0"))
        if lim:
            orders = orders[:lim]
        job["total"] = len(orders)

        def one(o):
            try:
                return o, waybill_one(o["id"]), None
            except Exception as e:
                return o, None, str(e)
        got = {}
        with ThreadPoolExecutor(4) as ex:
            for o, pdf, err in ex.map(one, orders):
                job["done"] += 1
                if err:
                    job["errors"].append(f"{o['code']}: {err}")
                else:
                    got[o["id"]] = pdf
        if not got:
            raise ValueError("ни одной накладной не получено")
        w = PdfWriter()
        for o in orders:  # порядок как в списке: одинаковые товары подряд
            if o["id"] in got:
                for pg in PdfReader(io.BytesIO(got[o["id"]])).pages:
                    w.add_page(pg)
        fn = f"waybills-{datetime.now(TZ):%Y-%m-%d_%H-%M}.pdf"
        with open(os.path.join(WB, fn), "wb") as fh:
            w.write(fh)
        log("Накладные", f"Сформировано {len(got)} из {len(orders)}, файл {fn}")
        with lock:
            db.executemany("delete from packed where oid=?", [(i,) for i in got])
            db.commit()
        job.update(state="done", file=fn)
    except Exception as e:
        job.update(state="error")
        job["errors"].append(str(e))
        log("Накладные", f"Ошибка: {e}")


ROUTES = {"/api/pack": lambda p: pack_view(), "/api/today": lambda p: today(), "/api/analytics": analytics, "/api/products": products,
          "/api/orders": orders, "/api/stock": lambda p: stock(), "/api/pricelist": pricelist, "/api/marketing": marketing, "/api/prefs": lambda p: prefs(), "/api/log": logs}


def backup_loop():
    while True:
        try:
            os.makedirs(BK, exist_ok=True)
            fn = os.path.join(BK, f"hub-{datetime.now(TZ):%Y-%m-%d}.db")
            if not os.path.exists(fn):
                dst = sqlite3.connect(fn)
                with lock:
                    db.backup(dst)
                dst.close()
                for old in sorted(glob.glob(os.path.join(BK, "hub-*.db")))[:-7]:
                    os.remove(old)
        except Exception as e:
            print("бэкап не удался:", e, flush=True)
        time.sleep(3600)


def migrate():
    # v2: раньше SKU позиций не сохранялся: перечитываем позиции заново
    if not q("select 1 from prefs where k='mig_entries2'"):
        with lock:
            db.execute("delete from entries")
            db.execute("update orders set items=0")
            db.execute("insert into prefs values('mig_entries2','1')")
            _log("Система", "Миграция: позиции заказов перечитываются заново (SKU)")
            db.commit()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def send(self, code, body, ctype="application/json"):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        if code == 401:
            self.send_header("WWW-Authenticate", 'Basic realm="Kaspi Hub"')
        self.end_headers()
        self.wfile.write(body)


    def authed(self):
        proxied = self.headers.get("CF-Connecting-IP") or self.headers.get("X-Forwarded-For")
        if not PASSWORD and not USERS:
            return not proxied  # без пароля пускаем только напрямую, через туннель/прокси нельзя
        ip = (proxied or self.client_address[0]).split(",")[-1].strip()
        now = time.time()
        f = [t for t in _fails.get(ip, []) if now - t < 600]
        _fails[ip] = f
        if len(f) >= 10:  # защита от подбора пароля: 10 неверных попыток за 10 минут
            return False
        h = self.headers.get("Authorization", "")
        if not h.startswith("Basic "):
            return False
        try:
            u, pw = base64.b64decode(h[6:]).decode().split(":", 1)
        except Exception:
            return False
        creds = ([(None, PASSWORD)] if PASSWORD else []) + USERS
        ok = False
        for n, p in creds:  # перебираем всех, чтобы время ответа не выдавало верный вариант
            ok |= hmac.compare_digest(pw.encode(), p.encode()) and (n is None or n == u)
        if not ok:
            f.append(now)
        return ok
    def do_POST(self):
        if not self.authed():
            return self.send(401, b'{"error":"auth"}')
        o = self.headers.get("Origin")
        if (o and urlparse(o).netloc != self.headers.get("Host")) or not self.headers.get("Content-Type", "").startswith("application/json"):
            return self.send(403, b'{"error":"forbidden"}')
        try:
            n = int(self.headers.get("Content-Length", 0))
            if n > (4_000_000 if self.path == "/api/stock_upload" else 100_000):
                raise ValueError("слишком большой запрос")
            self.send(200, json.dumps(post(urlparse(self.path).path, json.loads(self.rfile.read(n) or b"{}"))).encode())
        except (ValueError, KeyError, AssertionError, TypeError) as e:
            self.send(400, json.dumps({"error": "Некорректные данные" if not isinstance(e, ValueError) else str(e)}).encode())
        except Exception as e:
            print("POST ошибка:", e, flush=True)
            self.send(500, '{"error":"Внутренняя ошибка"}'.encode())

    def do_GET(self):
        if self.path == "/healthz":
            return self.send(200, b"ok", "text/plain")
        if not self.authed():
            return self.send(401, b"auth required", "text/plain")
        u = urlparse(self.path)
        try:
            if u.path in ROUTES:
                self.send(200, json.dumps({"data": ROUTES[u.path](parse_qs(u.query)), **info}).encode())
            elif re.fullmatch(r"/waybills/waybills-[\d_-]+\.pdf", u.path):
                fp = os.path.join(WB, os.path.basename(u.path))
                self.send(200, open(fp, "rb").read(), "application/pdf") if os.path.exists(fp) else self.send(404, b"not found", "text/plain")
            elif u.path in ("/", "/index.html"):
                self.send(200, open(os.path.join(DIR, "index.html"), "rb").read(), "text/html; charset=utf-8")
            else:
                self.send(404, b"not found", "text/plain")
        except Exception as e:
            print("GET ошибка:", u.path, e, flush=True)
            self.send(500, json.dumps({"data": None, "error": "Внутренняя ошибка сервера", "sync": info["sync"]}).encode())


for _c in ("sku text", "shows integer", "clicks integer", "orders integer"):
    try:
        db.execute(f"alter table ads add column {_c}")
    except sqlite3.OperationalError:
        pass
db.commit()


db.execute("create table if not exists packed(oid text primary key)")
db.commit()


if __name__ == "__main__":
    host = os.environ.get("HOST", "0.0.0.0" if os.environ.get("PORT") else "127.0.0.1")  # только этот компьютер: в данных есть имена клиентов
    if host not in ("127.0.0.1", "localhost") and not (PASSWORD or USERS):
        sys.exit("Для HOST вне localhost задайте пароль: HUB_PASSWORD=...")
    log("Система", "Сервер запущен")
    migrate()
    for t in (sync_loop, entries_loop, backup_loop):
        threading.Thread(target=t, daemon=True).start()
    print(f"Открой http://localhost:{os.environ.get('PORT', 8000)} (история за {BACKFILL} дн. грузится в фоне)")
    if host not in ("127.0.0.1", "localhost"):
        import socket
        try:
            sk = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            sk.connect(("8.8.8.8", 80))  # пакеты не отправляются, нужен только адрес этого компьютера в сети
            print(f"С телефона (тот же Wi-Fi): http://{sk.getsockname()[0]}:8000  пароль: из HUB_PASSWORD")
        except OSError:
            print("С телефона: http://<IP-этого-компьютера>:8000")
    ThreadingHTTPServer((host, int(os.environ.get("PORT", 8000))), Handler).serve_forever()
